"""Course fixtures with actual notes, reordered slides and a hidden slide."""
import zipfile
from xml.etree import ElementTree as ET

from pptx import Presentation
from pypdf import PdfWriter


def course_pptx(path, annotations=False):
    presentation = Presentation()
    for title, note in [("Introduction", "Notes de la première slide.\nDeuxième paragraphe."),
                        ("Slide masquée", "Notes de la slide masquée."),
                        ("Conclusion", "")]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "Texte du cours à lire pendant la création des cartes."
        slide.notes_slide.notes_text_frame.text = note
    presentation.slides[1]._element.set("show", "0")
    # Physical ZIP slide1/2/3 order differs from presentation display order.
    order = presentation.slides._sldIdLst
    order.insert(0, order[-1])
    presentation.save(path)
    if annotations:
        with zipfile.ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
        content_ns = "{http://schemas.openxmlformats.org/package/2006/content-types}"
        p = "http://schemas.openxmlformats.org/presentationml/2006/main"
        modern = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"
        drawing = "http://schemas.openxmlformats.org/drawingml/2006/main"
        types = ET.fromstring(parts["[Content_Types].xml"])

        def relation(part, identifier, kind, target):
            tree = ET.fromstring(parts[part])
            ET.SubElement(tree, rel_ns + "Relationship", Id=identifier, Type=kind, Target=target)
            parts[part] = ET.tostring(tree)

        def xml_part(name, content, content_type):
            parts[name] = content.encode()
            ET.SubElement(types, content_ns + "Override", PartName="/" + name, ContentType=content_type)

        # Valid package-relative metadata outside ppt/ reproduced the upload bug.
        relation("ppt/_rels/presentation.xml.rels", "rIdThumbnail", "http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail", "../docProps/thumbnail.jpeg")
        relation("ppt/_rels/presentation.xml.rels", "rIdClassicAuthors", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors", "commentAuthors.xml")
        relation("ppt/slides/_rels/slide1.xml.rels", "rIdClassicComments", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments", "../comments/comment1.xml")
        xml_part("ppt/commentAuthors.xml", f'<p:cmAuthorLst xmlns:p="{p}"><p:cmAuthor id="0" name="Professeur" initials="P" lastIdx="1"/></p:cmAuthorLst>', "application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml")
        xml_part("ppt/comments/comment1.xml", f'<p:cmLst xmlns:p="{p}"><p:cm authorId="0" idx="1" dt="2026-10-03T12:00:00Z"><p:pos x="0" y="0"/><p:text>Préciser cet exemple.</p:text></p:cm></p:cmLst>', "application/vnd.openxmlformats-officedocument.presentationml.comments+xml")
        relation("ppt/_rels/presentation.xml.rels", "rIdModernAuthors", "http://schemas.microsoft.com/office/2018/10/relationships/authors", "authors.xml")
        relation("ppt/slides/_rels/slide2.xml.rels", "rIdModernComments", "http://schemas.microsoft.com/office/2018/10/relationships/comments", "../comments/modern.xml")
        xml_part("ppt/authors.xml", f'<m:authorLst xmlns:m="{modern}"><m:author id="{{A1A11111-1111-4111-8111-111111111111}}" name="Enseignante" initials="E" userId="fixture" providerId=""/><m:author id="{{B2B22222-2222-4222-8222-222222222222}}" name="Étudiant" initials="É" userId="reply" providerId=""/></m:authorLst>', "application/vnd.ms-powerpoint.authors+xml")
        xml_part("ppt/comments/modern.xml", f'<m:cmLst xmlns:m="{modern}" xmlns:a="{drawing}"><m:cm id="{{D4D44444-4444-4444-8444-444444444444}}" authorId="{{A1A11111-1111-4111-8111-111111111111}}" created="2026-10-03T12:00:00Z"><m:unknownAnchor/><m:replyLst><m:reply id="{{C3C33333-3333-4333-8333-333333333333}}" authorId="{{B2B22222-2222-4222-8222-222222222222}}" created="2026-10-03T12:01:00Z"><m:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>Bien compris.</a:t></a:r></a:p></m:txBody></m:reply></m:replyLst><m:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>Commentaire sur la slide masquée.</a:t></a:r><a:br/><a:r><a:t>À revoir.</a:t></a:r></a:p></m:txBody></m:cm></m:cmLst>', "application/vnd.ms-powerpoint.comments+xml")
        slide = ET.fromstring(parts["ppt/slides/slide2.xml"])
        extensions = ET.SubElement(slide, "{" + p + "}extLst")
        extension = ET.SubElement(extensions, "{" + p + "}ext", uri="{6950BFC3-D8DA-4A85-94F7-54DA5524770B}")
        ET.SubElement(extension, "{" + modern + "}commentRel", {"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": "rIdModernComments"})
        parts["ppt/slides/slide2.xml"] = ET.tostring(slide)
        # LibreOffice's format detection requires the standard default namespace.
        ET.register_namespace("", content_ns[1:-1])
        parts["[Content_Types].xml"] = ET.tostring(types)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in parts.items():
                archive.writestr(name, content)
    return path


def course_pdf(path):
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=600, height=800)
    writer.write(path)
    return path
